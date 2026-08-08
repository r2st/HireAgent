"""Resume text extraction for PDF, DOCX, and plain text.

The design doc specifies Apache Tika; that needs a JVM sidecar, so this uses
pure-Python extractors (pypdf, python-docx) behind the same interface. Adding a
Tika backend later means adding one branch to ``extract_text`` — no caller
changes. Image-only resumes are detected and reported rather than silently
returning an empty parse.
"""

from __future__ import annotations

import io
import logging
import re

from app.core.errors import ValidationError

logger = logging.getLogger(__name__)

SUPPORTED_CONTENT_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "docx",
    "text/plain": "txt",
    "text/markdown": "txt",
}

SUPPORTED_EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".doc": "docx", ".txt": "txt", ".md": "txt"}

# Below this, a "successful" extraction is almost certainly a scanned image.
MIN_USEFUL_CHARS = 80

_WS_RUN = re.compile(r"[ \t ]+")
_BLANK_RUNS = re.compile(r"\n{3,}")


class ExtractionError(ValidationError):
    """Raised when a file cannot be turned into usable text."""


def detect_format(content_type: str | None, filename: str | None) -> str:
    """Resolve a file to ``pdf`` | ``docx`` | ``txt``.

    Content type is trusted first, then the extension, since browsers often
    send ``application/octet-stream`` for uploads.
    """
    if content_type:
        fmt = SUPPORTED_CONTENT_TYPES.get(content_type.split(";")[0].strip().lower())
        if fmt:
            return fmt
    if filename:
        lowered = filename.lower()
        for ext, fmt in SUPPORTED_EXTENSIONS.items():
            if lowered.endswith(ext):
                return fmt
    raise ExtractionError(
        "Unsupported resume format. Upload a PDF, DOCX, or plain-text file.",
        details={"content_type": content_type, "filename": filename},
    )


def clean_text(text: str) -> str:
    """Collapse extraction artefacts without destroying line structure.

    Line breaks are meaningful in resumes (they separate roles, bullets, and
    dates), so only horizontal whitespace and excess blank lines are collapsed.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Ligatures and control characters that survive PDF extraction.
    text = text.replace("ﬁ", "fi").replace("ﬂ", "fl").replace("•", "•")
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ch >= " ")
    text = _WS_RUN.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_RUNS.sub("\n\n", text).strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, ValueError, OSError) as exc:
        raise ExtractionError(f"Could not read the PDF: {exc}") from exc

    if reader.is_encrypted:
        # Some resumes are owner-password protected but readable with an empty
        # user password; try that before giving up.
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError(
                "The PDF is password protected and cannot be parsed."
            ) from exc

    pages: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:
            # One bad page should not lose the rest of the resume.
            logger.warning("Failed to extract PDF page %d: %s", index, exc)
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
        raise ExtractionError(f"Could not read the DOCX file: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    # Many resumes lay out experience in tables, which paragraphs miss entirely.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(
    data: bytes, *, content_type: str | None = None, filename: str | None = None
) -> tuple[str, str]:
    """Extract text from a resume file.

    Returns ``(text, detected_format)``. Raises ``ExtractionError`` when the
    format is unsupported, the file is unreadable, or the result is too short
    to be a real resume (typically a scanned image needing OCR).
    """
    if not data:
        raise ExtractionError("The uploaded file is empty")

    fmt = detect_format(content_type, filename)
    if fmt == "pdf":
        raw = _extract_pdf(data)
    elif fmt == "docx":
        raw = _extract_docx(data)
    else:
        raw = _extract_txt(data)

    text = clean_text(raw)
    if len(text) < MIN_USEFUL_CHARS:
        raise ExtractionError(
            "Could not extract readable text. The file may be a scanned image, "
            "which requires OCR.",
            details={"format": fmt, "extracted_chars": len(text)},
        )
    return text, fmt
