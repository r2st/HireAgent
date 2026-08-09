"""Synthetic resume files and fake LLM clients used across the test suite.

Resume fixtures are generated rather than checked in as binaries so the golden
inputs stay readable and reviewable in the diff.
"""

from __future__ import annotations

import io
from datetime import datetime
from types import ModuleType

import httpx
import pytest

from app.integrations.calendar import (
    BusyResult,
    CalendarEvent,
    CalendarProvider,
    EventResult,
)
from app.integrations.openrouter import LLMResult, OpenRouterClient


# --------------------------------------------------------------------------- #
# Fake HTTP transport
# --------------------------------------------------------------------------- #
class FakeResponse:
    def __init__(
        self, status_code: int, payload: object = None, text: str = ""
    ) -> None:
        self.status_code = status_code
        self._payload = payload
        self.text = text or str(payload)

    def json(self) -> object:
        if isinstance(self._payload, Exception):
            raise self._payload
        if self._payload is None:
            raise ValueError("no json body")
        return self._payload


class FakeHTTP:
    """Stands in for the ``httpx`` module inside a gateway.

    Responses are handed out in order, and every call is recorded so a test can
    assert on the request the gateway built, not merely on what it did with the
    reply. A queued ``Exception`` is raised instead of returned, which is how
    transport faults are simulated.
    """

    def __init__(self, *responses: object) -> None:
        self.responses = list(responses)
        self.calls: list[dict] = []
        outer = self

        class _Client:
            def __init__(self, **kwargs) -> None:
                self.kwargs = kwargs

            async def __aenter__(self):
                return self

            async def __aexit__(self, *exc) -> bool:
                return False

            async def _record(self, method: str, url: str, **kwargs):
                outer.calls.append({"method": method, "url": url, **kwargs})
                if not outer.responses:
                    raise AssertionError(f"unexpected {method} {url}")
                nxt = outer.responses.pop(0)
                if isinstance(nxt, Exception):
                    raise nxt
                return nxt

            async def post(self, url, **kwargs):
                return await self._record("POST", url, **kwargs)

            async def get(self, url, **kwargs):
                return await self._record("GET", url, **kwargs)

            async def delete(self, url, **kwargs):
                return await self._record("DELETE", url, **kwargs)

        self.AsyncClient = _Client
        self.HTTPError = httpx.HTTPError


def install_http(
    monkeypatch: pytest.MonkeyPatch, http: FakeHTTP, *modules: ModuleType
) -> FakeHTTP:
    """Swap ``httpx`` for ``http`` in each module that calls out.

    Every module doing its own request has to be named: token refresh lives in
    ``app.integrations.oauth`` while the API calls live in the gateway, so
    patching only the gateway lets a refresh reach the real network.
    """
    for module in modules:
        monkeypatch.setattr(module, "httpx", http)
    return http

# --------------------------------------------------------------------------- #
# Resume text
# --------------------------------------------------------------------------- #
SAMPLE_RESUME = """\
Ada Lovelace
Bengaluru, India
ada.lovelace@example.com | +91 98765 43210
https://linkedin.com/in/adalovelace
https://github.com/adalovelace
https://ada.dev

Professional Summary
Backend engineer with 8 years of professional experience building distributed
payment systems.

Technical Skills
Python, Django, PostgreSQL, Redis, Docker, Kubernetes, AWS, ReactJS

Work Experience
Staff Engineer at Analytical Engines | Jan 2021 - Present
Led the payments platform team.

Senior Engineer at Babbage Systems | Mar 2017 - Dec 2020
Built the ledger service.

Education
B.Tech Computer Science, Indian Institute of Technology, 2013 - 2017

Certifications
AWS Certified Solutions Architect
"""

# A deliberately sparse resume: no skills section, no dates, no phone.
MINIMAL_RESUME = """\
Grace Hopper
grace.hopper@example.com

Wrote a compiler. Interested in making computers easier to talk to. Currently
looking for a new role where I can keep doing that kind of work every day.
"""

# No email at all — ingestion must refuse to create a candidate from this.
ANONYMOUS_RESUME = """\
Confidential Candidate Profile

A senior backend engineer with deep experience in Python and PostgreSQL who
prefers to stay anonymous until the first conversation with the hiring team.
"""


# --------------------------------------------------------------------------- #
# File builders
# --------------------------------------------------------------------------- #
def make_pdf(text: str) -> bytes:
    """Build a minimal single-page PDF whose text pypdf can extract.

    Hand-rolled rather than pulled from a rendering library so the suite has no
    extra dependency for what is ultimately a few hundred bytes of fixture.
    """
    lines = text.split("\n")

    def escape(value: str) -> str:
        return value.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    ops = ["BT", "/F1 11 Tf", "13 TL", "40 780 Td"]
    for line in lines:
        ops.append(f"({escape(line)}) Tj")
        ops.append("T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]

    buffer = io.BytesIO()
    buffer.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(buffer.tell())
        buffer.write(f"{index} 0 obj\n".encode())
        buffer.write(body)
        buffer.write(b"\nendobj\n")

    xref_offset = buffer.tell()
    buffer.write(f"xref\n0 {len(objects) + 1}\n".encode())
    buffer.write(b"0000000000 65535 f \n")
    for offset in offsets:
        buffer.write(f"{offset:010d} 00000 n \n".encode())
    buffer.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode()
    )
    return buffer.getvalue()


def make_docx(text: str, *, table_rows: list[list[str]] | None = None) -> bytes:
    """Build a DOCX with one paragraph per line, plus an optional table."""
    import docx

    document = docx.Document()
    for line in text.split("\n"):
        document.add_paragraph(line)

    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for cell_index, value in enumerate(row):
                table.cell(row_index, cell_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Fake LLM clients
# --------------------------------------------------------------------------- #
class FakeLLMClient(OpenRouterClient):
    """An OpenRouter client that replays canned responses.

    ``responses`` is consumed one call at a time; the last entry repeats once
    exhausted so a test does not have to count how many calls a service makes.
    """

    def __init__(
        self,
        responses: list[dict | list | str | None] | None = None,
        *,
        available: bool = True,
        model: str = "fake/model:free",
    ) -> None:
        super().__init__(api_key="test-key", enabled=available)
        self.responses = list(responses or [])
        self.model = model
        self.calls: list[dict] = []

    @property
    def is_available(self) -> bool:
        return self._enabled

    def _next(self) -> dict | list | str | None:
        if not self.responses:
            return None
        if len(self.responses) == 1:
            return self.responses[0]
        return self.responses.pop(0)

    async def complete(self, **kwargs) -> LLMResult | None:
        self.calls.append(kwargs)
        payload = self._next()
        if payload is None:
            return None
        content = payload if isinstance(payload, str) else _dumps(payload)
        return LLMResult(content=content, model=self.model, latency_ms=7)

    async def complete_json(self, **kwargs) -> tuple[dict | list | None, LLMResult | None]:
        result = await self.complete(**kwargs)
        if result is None:
            return None, None
        from app.integrations.openrouter import extract_json

        return extract_json(result.content), result


def _dumps(payload: dict | list) -> str:
    import json

    return json.dumps(payload)


# --------------------------------------------------------------------------- #
# Fake calendar provider
# --------------------------------------------------------------------------- #
class FakeCalendarProvider(CalendarProvider):
    """A calendar that answers from a script instead of the network.

    ``busy`` is the block list every ``fetch_busy`` returns. Setting
    ``synced=False`` simulates an unreachable calendar, and ``raises`` makes the
    provider blow up so the callers' "a provider fault must not lose the
    booking" guards can be exercised.
    """

    name = "google"

    def __init__(
        self,
        busy: list[tuple[datetime, datetime]] | None = None,
        *,
        synced: bool = True,
        error: str | None = None,
        event_id: str = "ext-event-1",
        meeting_url: str = "https://meet.example.test/abc",
        create_ok: bool = True,
        delete_ok: bool = True,
        raises: bool = False,
    ) -> None:
        self.busy = list(busy or [])
        self.synced = synced
        self.error = error
        self.event_id = event_id
        self.meeting_url = meeting_url
        self.create_ok = create_ok
        self.delete_ok = delete_ok
        self.raises = raises
        self.created: list[CalendarEvent] = []
        self.deleted: list[str] = []
        self.busy_calls = 0

    @property
    def is_configured(self) -> bool:
        return True

    async def fetch_busy(self, credentials, start, end) -> BusyResult:
        self.busy_calls += 1
        if self.raises:
            raise RuntimeError("calendar exploded")
        return BusyResult(
            blocks=list(self.busy),
            synced=self.synced,
            error=self.error,
        )

    async def create_event(self, credentials, event: CalendarEvent) -> EventResult:
        if self.raises:
            raise RuntimeError("calendar exploded")
        self.created.append(event)
        if not self.create_ok:
            return EventResult(ok=False, error="calendar rejected the event")
        return EventResult(
            ok=True, external_event_id=self.event_id, meeting_url=self.meeting_url
        )

    async def delete_event(self, credentials, external_event_id: str) -> EventResult:
        if self.raises:
            raise RuntimeError("calendar exploded")
        self.deleted.append(external_event_id)
        if not self.delete_ok:
            return EventResult(ok=False, error="calendar refused the deletion")
        return EventResult(ok=True)


LLM_RESUME_PAYLOAD = {
    "full_name": "Ada Lovelace",
    "email": "ada.lovelace@example.com",
    "phone": "+91 98765 43210",
    "location": "Bengaluru, India",
    "linkedin_url": "https://linkedin.com/in/adalovelace",
    "github_url": "https://github.com/adalovelace",
    "portfolio_url": "https://ada.dev",
    "current_company": "Analytical Engines",
    "current_role": "Staff Engineer",
    "total_experience_years": 8,
    "summary": "Backend engineer building distributed payment systems.",
    "skills": [
        {"name": "python", "proficiency": "expert", "years": 8},
        {"name": "Django", "proficiency": "advanced", "years": None},
        {"name": "postgres", "proficiency": None, "years": 6},
        {"name": "ReactJS", "proficiency": "intermediate", "years": 2},
    ],
    "experience": [
        {
            "company": "Analytical Engines",
            "role": "Staff Engineer",
            "start_date": "2021-01",
            "end_date": None,
            "is_current": True,
            "location": "Bengaluru",
            "description": "Led the payments platform team.",
        },
        {
            "company": "Babbage Systems",
            "role": "Senior Engineer",
            "start_date": "2017-03",
            "end_date": "2020-12",
            "is_current": False,
            "location": "Pune",
            "description": "Built the ledger service.",
        },
    ],
    "education": [
        {
            "institution": "Indian Institute of Technology",
            "degree": "B.Tech",
            "field_of_study": "Computer Science",
            "start_year": 2013,
            "end_year": 2017,
            "grade": "8.6 CGPA",
        }
    ],
    "certifications": [
        {"name": "AWS Certified Solutions Architect", "issuer": "AWS", "year": 2022}
    ],
    "projects": [
        {
            "name": "Ledger",
            "description": "Double-entry accounting service.",
            "technologies": ["python", "postgres"],
        }
    ],
    "languages": ["English", "Hindi"],
}
